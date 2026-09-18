from datetime import datetime
from decimal import Decimal

from firebird_probe import connect
from local_report_export_common import (
    TEMPLATE_DIR,
    excel_value,
    output_path,
    period_label,
    run_cli,
    save_workbook,
)
from payment_deduplication import reportable_payment_filter


REPORT_NUMBER = 26


def clean(value):
    return str(value or "").strip()


def period_covered(years):
    values = sorted(year for year in years if year is not None)
    if not values:
        return ""
    if len(values) == 1:
        return str(values[0])
    if values == list(range(values[0], values[-1] + 1)):
        return f"{values[0]}-{values[-1]}"
    return ", ".join(str(year) for year in values)


def classification_label(description, code):
    code = clean(code)
    if code.upper().startswith("S"):
        return "SPECIAL"
    return clean(description) or code


def fetch_rows(date_from, date_to):
    report_year = datetime.strptime(date_from, "%Y-%m-%d").year
    paid_filter = reportable_payment_filter("p")
    sql = f"""
        SELECT
            p.PAYMENT_ID, p.PAYMENTDATE, p.PAIDBY, tx.OWNERNAME, p.RECEIPTNO,
            COALESCE(p.COLLECTOR, p.USERID), pcd.TAXTRANS_ID, pcd.ITAXTYPE_CT,
            pcd.CASETYPE_CT, pcd.TAXYEAR, pcd.AMOUNT, pcd.CLASSCODE_CT,
            pcd.PROPERTYKIND_CT, ra.TDNO, ra.TDNOFORGR, ra.PREDOMCLASSCODE_CT,
            prop.PINNO, prop.NEWPINNO, brgy.DESCRIPTION, cls.DESCRIPTION,
            kind.DESCRIPTION
        FROM PAYMENT p
        JOIN PAYMENTCLASSDETAIL pcd ON pcd.PAYMENT_ID = p.PAYMENT_ID
        LEFT JOIN TAXPAYER tx ON tx.LOCAL_TIN = p.LOCAL_TIN
        LEFT JOIN RPTASSESSMENT ra ON ra.TAXTRANS_ID = pcd.TAXTRANS_ID
        LEFT JOIN PROPERTY prop ON prop.PROP_ID = ra.PROP_ID
        LEFT JOIN T_BARANGAY brgy
               ON brgy.CODE = prop.BARANGAY_CT
              AND brgy.MUNICIPAL_ID = prop.MUNICIPAL_ID
              AND brgy.PROVINCE_CT = prop.PROVINCE_CT
        LEFT JOIN T_CLASSIFICATION cls
               ON cls.CODE = COALESCE(pcd.CLASSCODE_CT, ra.PREDOMCLASSCODE_CT)
        LEFT JOIN T_PROPERTYKIND kind
               ON kind.CODE = COALESCE(pcd.PROPERTYKIND_CT, prop.PROPERTYKIND_CT)
        WHERE p.PAYMENTDATE >= CAST(? AS DATE)
          AND p.PAYMENTDATE < DATEADD(1 DAY TO CAST(? AS DATE))
          AND p.PAYGROUP_CT = 'RPT'
          {paid_filter}
          AND COALESCE(pcd.CANCELLED_BV, 0) = 0
          AND pcd.TAXYEAR > ?
        ORDER BY p.PAYMENTDATE, p.RECEIPTNO, p.PAYMENT_ID, pcd.TAXTRANS_ID, pcd.TAXYEAR
    """
    records = {}
    order = []
    connection = connect()
    try:
        cursor = connection.cursor()
        cursor.execute(sql, (date_from, date_to, report_year))
        for row in cursor.fetchall():
            (
                payment_id, payment_date, paid_by, taxpayer_name, receipt_no,
                collector, taxtrans_id, tax_type, case_type, taxyear, amount,
                class_code, property_kind_code, td_no, td_no_for_gr,
                predom_class_code, pin_no, new_pin_no, barangay_name,
                classification_name, property_kind_name,
            ) = row
            classification_code = class_code or predom_class_code
            key = (payment_id, taxtrans_id, classification_label(None, classification_code))
            if key not in records:
                records[key] = {
                    "date": payment_date,
                    "paid_by": paid_by,
                    "taxpayer": taxpayer_name or paid_by,
                    "years": set(),
                    "pin": new_pin_no or pin_no,
                    "receipt": receipt_no,
                    "td": td_no_for_gr or td_no,
                    "barangay": barangay_name,
                    "bsc_gross": Decimal("0"),
                    "bsc_discount": Decimal("0"),
                    "sef_gross": Decimal("0"),
                    "sef_discount": Decimal("0"),
                    "classification": classification_label(classification_name, classification_code),
                    "property_kind": clean(property_kind_name) or clean(property_kind_code),
                    "collector": collector,
                }
                order.append(key)
            record = records[key]
            record["years"].add(taxyear)
            prefix = "bsc" if clean(tax_type) == "BSC" else "sef" if clean(tax_type) == "SEF" else None
            if prefix is None:
                continue
            if clean(case_type) == "DED":
                record[f"{prefix}_discount"] += abs(amount or Decimal("0"))
            else:
                record[f"{prefix}_gross"] += amount or Decimal("0")
        connection.rollback()
    finally:
        connection.close()

    rows = []
    for key in order:
        record = records[key]
        bsc_net = record["bsc_gross"] - record["bsc_discount"]
        sef_net = record["sef_gross"] - record["sef_discount"]
        rows.append([
            record["date"], record["paid_by"], record["taxpayer"], period_covered(record["years"]),
            record["pin"], record["receipt"], record["td"], record["barangay"],
            record["bsc_gross"], record["bsc_discount"], bsc_net,
            record["sef_gross"], record["sef_discount"], sef_net, bsc_net + sef_net,
            record["classification"], record["property_kind"], record["collector"],
        ])
    return rows


def export_report(date_from, date_to, output_dir):
    from openpyxl import load_workbook

    rows = fetch_rows(date_from, date_to)
    workbook = load_workbook(TEMPLATE_DIR / "RECORD OF REAL PROPERTY TAX COLLECTION - ADVANCE PAYMENT REPORT.xlsx")
    sheet = workbook.active
    sheet["E4"] = period_label(date_from, date_to)
    for row_index, values in enumerate(rows, start=11):
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column_index)
            cell.value = excel_value(value)
            if 9 <= column_index <= 15:
                cell.number_format = "#,##0.00"
    path = save_workbook(workbook, output_path(output_dir, REPORT_NUMBER, date_from, date_to))
    return path, len(rows)


if __name__ == "__main__":
    raise SystemExit(run_cli(REPORT_NUMBER, export_report))
