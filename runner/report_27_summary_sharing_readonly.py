from datetime import datetime
from decimal import Decimal

from local_report_export_common import (
    TEMPLATE_DIR,
    excel_value,
    output_path,
    run_cli,
    save_workbook,
)
from report_preview_readonly import build_report


REPORT_NUMBER = 27


def panel_row(panel, property_group, category):
    return next(
        (row for row in panel if row.get("property_group") == property_group and row.get("category") == category),
        {},
    )


def penalty_row(panel, property_group):
    current = panel_row(panel, property_group, "Current-Year Penalty")
    prior = panel_row(panel, property_group, "Prior-Year Penalty")
    return {
        key: Decimal(str(current.get(key) or 0)) + Decimal(str(prior.get(key) or 0))
        for key in ("amount", "provincial_share", "municipal_share", "barangay_share")
    }


def export_report(date_from, date_to, output_dir):
    from openpyxl import load_workbook

    payload = build_report(REPORT_NUMBER, date_from, date_to)
    workbook = load_workbook(TEMPLATE_DIR / "SUMMARY_REPORT_SHARING_TEMPLATE.xlsx")
    sheet = workbook.active
    start = datetime.strptime(date_from, "%Y-%m-%d")
    end = datetime.strptime(date_to, "%Y-%m-%d")
    sheet["F3"] = start.strftime("%B") if start.month == end.month else f"{start.strftime('%B')} to {end.strftime('%B')}"
    sheet["F4"] = str(start.day) if date_from == date_to else f"{start.day}-{end.day}"
    sheet["F5"] = str(end.year)

    for item in payload.get("template_cells", []):
        cell = sheet.cell(int(item["row"]), int(item["column"]))
        cell.value = excel_value(item.get("value"))
        cell.number_format = "#,##0.00"

    row_map = {
        ("Land", "Current"): 36, ("Land", "Prior"): 37,
        ("Land", "Penalties"): 38, ("Land", "TOTAL"): 39,
        ("Building", "Current"): 46, ("Building", "Prior"): 47,
        ("Building", "Penalties"): 48, ("Building", "TOTAL"): 49,
    }
    for tax_type, columns in (("BSC", (3, 4, 5, 6)), ("SEF", (10, 11, 12))):
        panel = payload["share_panels"][tax_type]
        for property_group in ("Land", "Building"):
            values_by_category = {
                "Current": panel_row(panel, property_group, "Current"),
                "Prior": panel_row(panel, property_group, "Prior"),
                "Penalties": penalty_row(panel, property_group),
                "TOTAL": panel_row(panel, property_group, "TOTAL"),
            }
            for category, row in values_by_category.items():
                values = [row.get("amount", 0), row.get("provincial_share", 0), row.get("municipal_share", 0)]
                if tax_type == "BSC":
                    values.append(row.get("barangay_share", 0))
                for column, value in zip(columns, values):
                    cell = sheet.cell(row_map[(property_group, category)], column)
                    cell.value = excel_value(value)
                    cell.number_format = "#,##0.00"

    sheet["C51"] = "=SUM(D49:F49)"
    sheet["C45"] = "BUILDING"
    sheet["J45"] = "BUILDING"
    sheet["B26"] = "BLDG-INDUS/SPECIAL"
    sheet["I26"] = "BLDG-INDUS/SPECIAL"
    sheet["I51"] = "BUILDING SHARING TOTAL"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    path = save_workbook(workbook, output_path(output_dir, REPORT_NUMBER, date_from, date_to))
    return path, len(payload.get("template_cells", []))


if __name__ == "__main__":
    raise SystemExit(run_cli(REPORT_NUMBER, export_report))
