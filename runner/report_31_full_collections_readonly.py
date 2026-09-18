from datetime import datetime

from local_report_export_common import (
    TEMPLATE_DIR,
    excel_value,
    output_path,
    run_cli,
    save_workbook,
)
from report_preview_readonly import build_report


REPORT_NUMBER = 31


def export_report(date_from, date_to, output_dir):
    from openpyxl import load_workbook

    payload = build_report(REPORT_NUMBER, date_from, date_to)
    rows = payload.get("rows", [])
    workbook = load_workbook(TEMPLATE_DIR / "FULL_REPORT_COLLECTIONS.xlsx")
    sheet = workbook.active
    start = datetime.strptime(date_from, "%Y-%m-%d")
    end = datetime.strptime(date_to, "%Y-%m-%d")
    sheet["D4"] = start.strftime("%B") if start.month == end.month else f"{start.strftime('%B')} to {end.strftime('%B')}"
    sheet["D5"] = str(end.year)

    first_row = 8
    template_capacity = 24
    if len(rows) > template_capacity:
        sheet.insert_rows(32, len(rows) - template_capacity)
    total_row = first_row + max(len(rows), template_capacity)

    for row_index in range(first_row, total_row):
        for column_index in range(1, 7):
            sheet.cell(row_index, column_index).value = None
    for row_index, row in enumerate(rows, start=first_row):
        values = (row.get("date"), row.get("ctc"), row.get("rpt"), row.get("gf_tf"))
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column_index)
            cell.value = excel_value(value)
            if column_index > 1:
                cell.number_format = "#,##0.00"
        sheet.cell(row_index, 5).value = None
        sheet.cell(row_index, 6).value = f"=SUM(B{row_index}:D{row_index})"
        sheet.cell(row_index, 6).number_format = "#,##0.00"

    sheet.cell(total_row, 1).value = "TOTAL"
    for column in "BCDEF":
        sheet[f"{column}{total_row}"] = f"=SUM({column}{first_row}:{column}{total_row - 1})"
        sheet[f"{column}{total_row}"].number_format = "#,##0.00"
    sheet[f"C{total_row + 4}"] = "RCD TOTAL"
    sheet[f"F{total_row + 4}"] = f"=F{total_row}"
    sheet[f"C{total_row + 5}"] = "LESS: DUE FROM"
    sheet[f"F{total_row + 5}"] = f"=E{total_row}"
    sheet[f"C{total_row + 6}"] = "TOTAL COLLECTIONS"
    sheet[f"F{total_row + 6}"] = f"=F{total_row + 4}-F{total_row + 5}"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    path = save_workbook(workbook, output_path(output_dir, REPORT_NUMBER, date_from, date_to))
    return path, len(rows)


if __name__ == "__main__":
    raise SystemExit(run_cli(REPORT_NUMBER, export_report))
