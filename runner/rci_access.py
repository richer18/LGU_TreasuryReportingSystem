#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from copy import copy

import pyodbc
from openpyxl import load_workbook

DRIVER = "Microsoft Access Driver (*.mdb, *.accdb)"
TABLE = "RCI_CHECKS"
AUDIT = "RCI_CHECK_AUDIT"

FIELDS = [
    "CheckDate", "CheckNumber", "DVNumber", "FundType", "BankName", "BankAccountNumber",
    "Payee", "NatureOfPayment", "DVAmount", "ReportNumber", "SheetNumber", "ReportingMonth",
    "ReportingYear", "Status", "Remarks",
]

SORT_MAP = {
    "check_date": "CheckDate",
    "check_number": "CheckNumber",
    "dv_number": "DVNumber",
    "fund_type": "FundType",
    "bank_name": "BankName",
    "payee": "Payee",
    "dv_amount": "DVAmount",
    "status": "Status",
    "report_number": "ReportNumber",
}


def emit(payload: dict, code: int = 0) -> None:
    print(json.dumps(payload, ensure_ascii=False, default=str))
    raise SystemExit(code)


def db_path() -> Path:
    configured = os.environ.get("RCI_ACCESS_DB")
    if configured:
        return Path(configured)
    return Path(__file__).resolve().parents[1] / "backend" / "database" / "rci" / "LGU_RCI.accdb"


def template_path() -> Path:
    configured = os.environ.get("RCI_TEMPLATE_PATH")
    if configured:
        return Path(configured)
    return Path(r"L:\RCI\RCI- Template.xlsx")


def export_dir() -> Path:
    path = Path(__file__).resolve().parents[1] / "backend" / "storage" / "app" / "rci-exports"
    path.mkdir(parents=True, exist_ok=True)
    return path


def connect():
    path = db_path()
    if not path.exists():
        raise RuntimeError(f"RCI Access database was not found: {path}")
    return pyodbc.connect(f"DRIVER={{{DRIVER}}};DBQ={path};", autocommit=False)


def clean(value):
    value = "" if value is None else str(value).strip()
    return value or None


def to_money(value) -> str:
    if value in (None, ""):
        return "0.00"
    try:
        amount = Decimal(str(value).replace(",", ""))
    except (InvalidOperation, ValueError):
        raise ValueError("DV amount must be numeric.")
    if amount < 0:
        raise ValueError("DV amount must be greater than or equal to zero.")
    return f"{amount:.2f}"


def parse_date(value):
    value = clean(value)
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            pass
    raise ValueError("Check date must be valid.")


def date_out(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value)


def dt_out(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def ensure_schema(conn):
    cur = conn.cursor()
    tables = {row.table_name.upper() for row in cur.tables(tableType="TABLE")}
    if TABLE not in tables:
        cur.execute(f"""
            CREATE TABLE {TABLE} (
                ID AUTOINCREMENT PRIMARY KEY,
                CheckDate DATETIME,
                CheckNumber TEXT(80),
                DVNumber TEXT(120),
                FundType TEXT(120),
                BankName TEXT(120),
                BankAccountNumber TEXT(120),
                Payee TEXT(255),
                NatureOfPayment LONGTEXT,
                DVAmount CURRENCY,
                ReportNumber TEXT(120),
                SheetNumber INTEGER,
                ReportingMonth BYTE,
                ReportingYear SMALLINT,
                Status TEXT(30),
                Remarks LONGTEXT,
                CreatedBy INTEGER,
                UpdatedBy INTEGER,
                CancelledBy INTEGER,
                CancelledAt DATETIME,
                VoidedBy INTEGER,
                VoidedAt DATETIME,
                DeletedBy INTEGER,
                DeletedAt DATETIME,
                RestoredBy INTEGER,
                CreatedAt DATETIME,
                UpdatedAt DATETIME
            )
        """)
    if AUDIT not in tables:
        cur.execute(f"""
            CREATE TABLE {AUDIT} (
                ID AUTOINCREMENT PRIMARY KEY,
                CheckRecordID INTEGER,
                ActionName TEXT(40),
                OldValues LONGTEXT,
                NewValues LONGTEXT,
                PerformedBy INTEGER,
                PerformedAt DATETIME
            )
        """)
    for name, sql in {
        "ux_rci_check_number": f"CREATE UNIQUE INDEX ux_rci_check_number ON {TABLE} (CheckNumber)",
        "ux_rci_dv_number": f"CREATE UNIQUE INDEX ux_rci_dv_number ON {TABLE} (DVNumber)",
        "ix_rci_period": f"CREATE INDEX ix_rci_period ON {TABLE} (ReportingYear, ReportingMonth)",
        "ix_rci_status": f"CREATE INDEX ix_rci_status ON {TABLE} (Status)",
    }.items():
        try:
            cur.execute(sql)
        except pyodbc.Error:
            pass
    conn.commit()


def normalize_payload(payload):
    check_date = parse_date(payload.get("check_date"))
    return {
        "CheckDate": check_date,
        "CheckNumber": clean(payload.get("check_number")),
        "DVNumber": clean(payload.get("dv_number")),
        "FundType": clean(payload.get("fund_type")),
        "BankName": clean(payload.get("bank_name")),
        "BankAccountNumber": clean(payload.get("bank_account_number")),
        "Payee": clean(payload.get("payee")),
        "NatureOfPayment": clean(payload.get("nature_of_payment")),
        "DVAmount": to_money(payload.get("dv_amount")),
        "ReportNumber": clean(payload.get("report_number")),
        "SheetNumber": int(payload["sheet_number"]) if clean(payload.get("sheet_number")) else None,
        "ReportingMonth": int(payload["reporting_month"]) if clean(payload.get("reporting_month")) else (check_date.month if check_date else None),
        "ReportingYear": int(payload["reporting_year"]) if clean(payload.get("reporting_year")) else (check_date.year if check_date else None),
        "Status": clean(payload.get("status")) or "Draft",
        "Remarks": clean(payload.get("remarks")),
        "CreatedBy": int(payload["created_by"]) if clean(payload.get("created_by")) else None,
        "UpdatedBy": int(payload["updated_by"]) if clean(payload.get("updated_by")) else None,
    }


def row_to_record(row):
    return {
        "id": row.ID,
        "check_date": date_out(row.CheckDate),
        "check_number": row.CheckNumber,
        "dv_number": row.DVNumber,
        "fund_type": row.FundType,
        "bank_name": row.BankName,
        "bank_account_number": row.BankAccountNumber,
        "payee": row.Payee,
        "nature_of_payment": row.NatureOfPayment,
        "dv_amount": f"{Decimal(str(row.DVAmount or 0)):.2f}",
        "report_number": row.ReportNumber,
        "sheet_number": row.SheetNumber,
        "reporting_month": row.ReportingMonth,
        "reporting_year": row.ReportingYear,
        "status": row.Status or "Draft",
        "remarks": row.Remarks,
        "created_by": row.CreatedBy,
        "updated_by": row.UpdatedBy,
        "cancelled_by": row.CancelledBy,
        "cancelled_at": dt_out(row.CancelledAt),
        "voided_by": row.VoidedBy,
        "voided_at": dt_out(row.VoidedAt),
        "deleted_by": row.DeletedBy,
        "created_at": dt_out(row.CreatedAt),
        "updated_at": dt_out(row.UpdatedAt),
    }


def load_payload(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def fetch_all(conn, filters, issued_only=False, include_deleted=False):
    where = []
    params = []
    if not include_deleted:
        where.append("DeletedAt IS NULL")
    if issued_only:
        where.append("Status = ?")
        params.append("Issued")
    simple = {
        "reporting_month": ("ReportingMonth", int),
        "reporting_year": ("ReportingYear", int),
        "fund_type": ("FundType", str),
        "status": ("Status", str),
    }
    for key, (column, caster) in simple.items():
        value = clean(filters.get(key))
        if value:
            where.append(f"{column} = ?")
            params.append(caster(value))
    for key, column in {
        "bank_name": "BankName",
        "payee": "Payee",
        "check_number": "CheckNumber",
        "dv_number": "DVNumber",
    }.items():
        value = clean(filters.get(key))
        if value:
            where.append(f"{column} LIKE ?")
            params.append(f"%{value}%")
    search = clean(filters.get("search"))
    if search:
        where.append("(CheckNumber LIKE ? OR DVNumber LIKE ? OR Payee LIKE ? OR NatureOfPayment LIKE ? OR BankName LIKE ? OR ReportNumber LIKE ?)")
        params.extend([f"%{search}%"] * 6)
    sort_by = SORT_MAP.get(clean(filters.get("sort_by")) or "check_date", "CheckDate")
    sort_dir = "ASC" if clean(filters.get("sort_dir")) == "asc" else "DESC"
    sql = f"SELECT * FROM {TABLE}"
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += f" ORDER BY {sort_by} {sort_dir}, ID DESC"
    return [row_to_record(row) for row in conn.cursor().execute(sql, params).fetchall()]


def is_incomplete(row):
    if row.get("status") == "Draft":
        return True
    required = ["check_date", "check_number", "dv_number", "fund_type", "payee", "nature_of_payment"]
    return any(not row.get(field) for field in required) or Decimal(str(row.get("dv_amount") or 0)) <= 0


def duplicate_groups(records):
    groups = {}
    for row in records:
        if not row.get("payee") or not row.get("check_date") or not row.get("nature_of_payment"):
            continue
        amount = Decimal(str(row.get("dv_amount") or 0))
        if amount <= 0:
            continue
        key = (row["payee"].strip().lower(), row["check_date"], f"{amount:.2f}", row["nature_of_payment"].strip().lower())
        groups.setdefault(key, []).append(row)
    return [items for items in groups.values() if len(items) > 1]


def numeric_tail(value):
    if not value:
        return None
    digits = ""
    for char in reversed(str(value)):
        if char.isdigit():
            digits = char + digits
        elif digits:
            break
    return int(digits) if digits else None


def build_summary(records):
    issued = [row for row in records if row.get("status") == "Issued"]
    subtotals = {}
    for row in issued:
        fund = row.get("fund_type") or "Unspecified"
        item = subtotals.setdefault(fund, {"fund_type": fund, "issued_checks": 0, "issued_amount": Decimal("0.00")})
        item["issued_checks"] += 1
        item["issued_amount"] += Decimal(str(row.get("dv_amount") or 0))
    total = sum((Decimal(str(row.get("dv_amount") or 0)) for row in issued), Decimal("0.00"))
    return {
        "total_records": len(records),
        "total_issued_amount": f"{total:.2f}",
        "total_issued_checks": len(issued),
        "incomplete_records": sum(1 for row in records if is_incomplete(row)),
        "voided_checks": sum(1 for row in records if row.get("status") == "Voided"),
        "cancelled_checks": sum(1 for row in records if row.get("status") == "Cancelled"),
        "possible_duplicates": len(duplicate_groups(records)),
        "fund_subtotals": [
            {**item, "issued_amount": f"{item['issued_amount']:.2f}"}
            for item in subtotals.values()
        ],
        "grand_total": f"{total:.2f}",
    }


def build_warnings(records):
    warnings = []
    for group in duplicate_groups(records):
        warnings.append({
            "type": "possible_duplicate",
            "message": "Possible duplicate check based on payee, amount, date, and nature of payment.",
            "record_ids": [row["id"] for row in group],
        })
    numbers = sorted(
        [(row["id"], numeric_tail(row.get("check_number"))) for row in records if row.get("status") == "Issued"],
        key=lambda item: item[1] if item[1] is not None else -1,
    )
    numbers = [item for item in numbers if item[1] is not None]
    for previous, current in zip(numbers, numbers[1:]):
        if current[1] > previous[1] + 1:
            warnings.append({
                "type": "sequence_gap",
                "message": f"Check number sequence gap detected between {previous[1]} and {current[1]}.",
                "record_ids": [previous[0], current[0]],
            })
    return warnings


def audit(cur, record_id, action, old, new, user_id):
    cur.execute(
        f"INSERT INTO {AUDIT} (CheckRecordID, ActionName, OldValues, NewValues, PerformedBy, PerformedAt) VALUES (?, ?, ?, ?, ?, ?)",
        record_id,
        action,
        json.dumps(old, ensure_ascii=False, default=str) if old is not None else None,
        json.dumps(new, ensure_ascii=False, default=str) if new is not None else None,
        user_id,
        datetime.now(),
    )


def duplicate_exists(cur, check_number, dv_number, ignore_id=None):
    checks = [("CheckNumber", check_number, "Check number already exists."), ("DVNumber", dv_number, "DV number already exists.")]
    for column, value, message in checks:
        if not value:
            continue
        sql = f"SELECT ID FROM {TABLE} WHERE {column} = ? AND DeletedAt IS NULL"
        params = [value]
        if ignore_id:
            sql += " AND ID <> ?"
            params.append(int(ignore_id))
        if cur.execute(sql, params).fetchone():
            emit({"ok": False, "duplicate": True, "error": message}, 2)


def get_record(cur, record_id, include_deleted=False):
    sql = f"SELECT * FROM {TABLE} WHERE ID = ?"
    params = [int(record_id)]
    if not include_deleted:
        sql += " AND DeletedAt IS NULL"
    row = cur.execute(sql, params).fetchone()
    if not row:
        raise RuntimeError("Check record was not found.")
    return row_to_record(row)


def list_records(filters):
    with connect() as conn:
        ensure_schema(conn)
        records = fetch_all(conn, filters)
        total = len(records)
        per_page = max(5, min(100, int(filters.get("per_page") or 15)))
        page = max(1, int(filters.get("page") or 1))
        start = (page - 1) * per_page
        page_rows = records[start:start + per_page]
        return {
            "records": page_rows,
            "meta": {
                "current_page": page,
                "last_page": max(1, (total + per_page - 1) // per_page),
                "per_page": per_page,
                "total": total,
            },
            "summary": build_summary(records),
            "warnings": build_warnings(records),
        }


def store_record(payload):
    data = normalize_payload(payload)
    now = datetime.now()
    user_id = data.pop("UpdatedBy", None) or data.get("CreatedBy")
    with connect() as conn:
        ensure_schema(conn)
        cur = conn.cursor()
        duplicate_exists(cur, data["CheckNumber"], data["DVNumber"])
        if data["Status"] == "Cancelled":
            data["CancelledBy"] = user_id
            data["CancelledAt"] = now
        if data["Status"] == "Voided":
            data["VoidedBy"] = user_id
            data["VoidedAt"] = now
        data["UpdatedBy"] = user_id
        data["CreatedAt"] = now
        data["UpdatedAt"] = now
        columns = list(data.keys())
        cur.execute(
            f"INSERT INTO {TABLE} ({', '.join(f'[{column}]' for column in columns)}) VALUES ({', '.join('?' for _ in columns)})",
            [data[column] for column in columns],
        )
        record_id = int(cur.execute("SELECT @@IDENTITY").fetchone()[0])
        record = get_record(cur, record_id)
        audit(cur, record_id, "created", None, record, user_id)
        conn.commit()
        return record


def update_record(payload):
    record_id = int(payload["id"])
    data = normalize_payload(payload)
    data.pop("CreatedBy", None)
    user_id = data.get("UpdatedBy")
    with connect() as conn:
        ensure_schema(conn)
        cur = conn.cursor()
        old = get_record(cur, record_id)
        duplicate_exists(cur, data["CheckNumber"], data["DVNumber"], record_id)
        if data["Status"] == "Cancelled" and not old.get("cancelled_at"):
            data["CancelledBy"] = user_id
            data["CancelledAt"] = datetime.now()
        if data["Status"] == "Voided" and not old.get("voided_at"):
            data["VoidedBy"] = user_id
            data["VoidedAt"] = datetime.now()
        data["UpdatedAt"] = datetime.now()
        assignments = ", ".join(f"[{column}] = ?" for column in data.keys())
        cur.execute(f"UPDATE {TABLE} SET {assignments} WHERE ID = ?", [*data.values(), record_id])
        new = get_record(cur, record_id)
        audit(cur, record_id, "updated", old, new, user_id)
        conn.commit()
        return new


def status_record(payload):
    record_id = int(payload["id"])
    status = clean(payload.get("status")) or "Draft"
    remarks = clean(payload.get("remarks"))
    user_id = int(payload["updated_by"]) if clean(payload.get("updated_by")) else None
    with connect() as conn:
        ensure_schema(conn)
        cur = conn.cursor()
        old = get_record(cur, record_id)
        fields = {"Status": status, "Remarks": remarks or old.get("remarks"), "UpdatedBy": user_id, "UpdatedAt": datetime.now()}
        if status == "Cancelled" and not old.get("cancelled_at"):
            fields["CancelledBy"] = user_id
            fields["CancelledAt"] = datetime.now()
        if status == "Voided" and not old.get("voided_at"):
            fields["VoidedBy"] = user_id
            fields["VoidedAt"] = datetime.now()
        assignments = ", ".join(f"[{column}] = ?" for column in fields.keys())
        cur.execute(f"UPDATE {TABLE} SET {assignments} WHERE ID = ?", [*fields.values(), record_id])
        new = get_record(cur, record_id)
        audit(cur, record_id, status.lower(), old, new, user_id)
        conn.commit()
        return new


def delete_record(record_id, user_id):
    with connect() as conn:
        ensure_schema(conn)
        cur = conn.cursor()
        old = get_record(cur, record_id)
        now = datetime.now()
        cur.execute(f"UPDATE {TABLE} SET DeletedBy = ?, DeletedAt = ?, UpdatedAt = ? WHERE ID = ?", user_id, now, now, int(record_id))
        new = get_record(cur, record_id, include_deleted=True)
        audit(cur, int(record_id), "deleted", old, new, user_id)
        conn.commit()


def restore_record(record_id, user_id):
    with connect() as conn:
        ensure_schema(conn)
        cur = conn.cursor()
        old = get_record(cur, record_id, include_deleted=True)
        now = datetime.now()
        cur.execute(f"UPDATE {TABLE} SET DeletedBy = NULL, DeletedAt = NULL, RestoredBy = ?, UpdatedAt = ? WHERE ID = ?", user_id, now, int(record_id))
        new = get_record(cur, record_id)
        audit(cur, int(record_id), "restored", old, new, user_id)
        conn.commit()
        return new


def show_record(record_id):
    with connect() as conn:
        ensure_schema(conn)
        cur = conn.cursor()
        record = get_record(cur, record_id)
        audits = []
        for row in cur.execute(f"SELECT * FROM {AUDIT} WHERE CheckRecordID = ? ORDER BY PerformedAt DESC, ID DESC", int(record_id)).fetchall():
            audits.append({
                "id": row.ID,
                "action": row.ActionName,
                "old_values": json.loads(row.OldValues) if row.OldValues else None,
                "new_values": json.loads(row.NewValues) if row.NewValues else None,
                "performed_by": row.PerformedBy,
                "performed_at": dt_out(row.PerformedAt),
            })
        return {"record": record, "audit_trail": audits}


def month_name(month):
    return datetime(2026, int(month or 1), 1).strftime("%B").upper()


def copy_row_format(ws, source_row, target_row):
    for col in range(1, ws.max_column + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.font:
            target.font = copy(source.font)
        if source.fill:
            target.fill = copy(source.fill)
        if source.border:
            target.border = copy(source.border)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def fund_label(fund):
    labels = {
        "General Fund - LBP": "GENERAL FUND (GF) - LBP",
        "Trust Fund Regular": "TRUST FUND REG. ",
        "Veterans": "VETERANS ",
        "DBP": "DBP",
        "SEF": "SEF",
        "Trust Liability DRRM": "TF - STF (TRUST LIABILITY DRRM)",
        "KALAHI": "TF - KCNCDDP (KALAHI)",
        "BUB": "BUB",
        "EGOV": "EGOV",
    }
    return labels.get(fund, fund or "Unspecified")


def find_fund_blocks(ws):
    known = {
        "GENERAL FUND (GF) - LBP",
        "TRUST FUND REG.",
        "VETERANS",
        "DBP",
        "SEF",
        "TF - STF (TRUST LIABILITY DRRM)",
        "TF - KCNCDDP (KALAHI)",
        "BUB",
        "EGOV",
    }
    blocks = []
    for row in range(1, ws.max_row + 1):
        value = str(ws.cell(row, 1).value or "").strip()
        if value in known:
            blocks.append({"label": value, "title_row": row, "header_row": row + 2, "first_data_row": row + 7 if value == "GENERAL FUND (GF) - LBP" else row + 3})
    for index, block in enumerate(blocks):
        next_title = blocks[index + 1]["title_row"] if index + 1 < len(blocks) else ws.max_row + 1
        follows = None
        for row in range(block["first_data_row"], next_title):
            if "nothing follows" in str(ws.cell(row, 4).value or "").lower():
                follows = row
                break
        block["follows_row"] = follows or max(block["first_data_row"], next_title - 1)
    return blocks


def clear_block(ws, block):
    for row in range(block["first_data_row"], block["follows_row"]):
        for col in range(1, 7):
            ws.cell(row, col).value = None


def write_block(ws, block, rows):
    clear_block(ws, block)
    needed = max(len(rows), 1)
    available = max(1, block["follows_row"] - block["first_data_row"])
    if needed > available:
        insert_at = block["follows_row"]
        ws.insert_rows(insert_at, needed - available)
        for row in range(insert_at, insert_at + needed - available):
            copy_row_format(ws, block["first_data_row"], row)
        block["follows_row"] += needed - available
    for offset, row in enumerate(rows):
        target = block["first_data_row"] + offset
        ws.cell(target, 1).value = row.get("check_date")
        ws.cell(target, 2).value = row.get("check_number")
        ws.cell(target, 3).value = row.get("dv_number")
        ws.cell(target, 4).value = row.get("payee")
        ws.cell(target, 5).value = row.get("nature_of_payment")
        ws.cell(target, 6).value = float(row.get("dv_amount") or 0)
    ws.cell(block["follows_row"], 4).value = "_ nothing follows _"
    if rows:
        ws.cell(block["follows_row"], 6).value = f"=SUM(F{block['first_data_row']}:F{block['first_data_row'] + len(rows) - 1})"
    else:
        ws.cell(block["follows_row"], 6).value = None


def export_xlsx(filters):
    path = template_path()
    if not path.exists():
        raise RuntimeError(f"RCI template was not found: {path}")
    with connect() as conn:
        ensure_schema(conn)
        rows = fetch_all(conn, filters, issued_only=True)
    wb = load_workbook(path)
    ws = wb.active
    month = int(filters.get("reporting_month") or datetime.now().month)
    year = int(filters.get("reporting_year") or datetime.now().year)
    ws["A5"] = f"{month_name(month)} 1-31, {year}"
    ws["A6"] = "Period Covered"
    grouped = {}
    for row in rows:
        grouped.setdefault(row.get("fund_type") or "Unspecified", []).append(row)
    blocks = find_fund_blocks(ws)
    for block in blocks:
        title = str(ws.cell(block["title_row"], 1).value or "").strip()
        fund_rows = []
        for fund, items in grouped.items():
            if fund_label(fund).strip() == title:
                fund_rows = items
                break
        if fund_rows:
            first = fund_rows[0]
            ws.cell(block["title_row"] + 1, 6).value = first.get("report_number")
            ws.cell(block["title_row"] + 2, 6).value = first.get("sheet_number")
            if first.get("bank_name") or first.get("bank_account_number"):
                ws.cell(block["title_row"] + 2, 1).value = f"Bank Name/Account No.    {first.get('bank_name') or ''}/{first.get('bank_account_number') or ''}"
        write_block(ws, block, fund_rows)
    filename = f"report_of_checks_issued_{year}_{month:02d}_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    output = export_dir() / filename
    wb.save(output)
    return {"path": str(output), "filename": filename, "records": rows}


def main():
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    for name in ("list", "store", "update", "status", "export", "export-xlsx"):
        cmd = sub.add_parser(name)
        cmd.add_argument("--payload-file", required=True)
    show = sub.add_parser("show")
    show.add_argument("--id", required=True)
    delete = sub.add_parser("delete")
    delete.add_argument("--id", required=True)
    delete.add_argument("--user-id", default="")
    restore = sub.add_parser("restore")
    restore.add_argument("--id", required=True)
    restore.add_argument("--user-id", default="")
    sub.add_parser("ensure")
    args = parser.parse_args()

    try:
        if args.command == "ensure":
            with connect() as conn:
                ensure_schema(conn)
            emit({"ok": True, "data": {"database": str(db_path())}})
        if args.command == "list":
            emit({"ok": True, "data": list_records(load_payload(args.payload_file))})
        if args.command == "export":
            filters = load_payload(args.payload_file)
            with connect() as conn:
                ensure_schema(conn)
                emit({"ok": True, "data": {"records": fetch_all(conn, filters, issued_only=True)}})
        if args.command == "export-xlsx":
            emit({"ok": True, "data": export_xlsx(load_payload(args.payload_file))})
        if args.command == "store":
            emit({"ok": True, "data": store_record(load_payload(args.payload_file))})
        if args.command == "show":
            emit({"ok": True, "data": show_record(args.id)})
        if args.command == "update":
            emit({"ok": True, "data": update_record(load_payload(args.payload_file))})
        if args.command == "status":
            emit({"ok": True, "data": status_record(load_payload(args.payload_file))})
        if args.command == "delete":
            delete_record(args.id, int(args.user_id) if clean(args.user_id) else None)
            emit({"ok": True, "data": {}})
        if args.command == "restore":
            emit({"ok": True, "data": restore_record(args.id, int(args.user_id) if clean(args.user_id) else None)})
    except SystemExit:
        raise
    except Exception as exc:
        emit({"ok": False, "error": str(exc)}, 1)


if __name__ == "__main__":
    main()
