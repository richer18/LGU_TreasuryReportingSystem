from datetime import datetime
from decimal import Decimal

from business_permit_report_readonly import build_payload
from local_report_export_common import (
    TEMPLATE_DIR,
    excel_value,
    output_path,
    run_cli,
    save_workbook,
)


REPORT_NUMBER = 32


def clean(value):
    return " ".join(str(value or "").replace("\n", " ").split()).strip()


def parse_date(value):
    try:
        return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def normalize_business_type(value):
    text = clean(value).upper()
    if "ONE" in text and "CORPORATION" in text:
        return "One-Person Corporation"
    if "CORPORATION" in text:
        return "Corporation"
    if "PARTNERSHIP" in text:
        return "Partnership"
    if "COOPERATIVE" in text:
        return "Cooperative"
    return "Single Proprietor"


def capitalization_size(capital, gross):
    basis = Decimal(str(capital or gross or 0))
    if basis <= Decimal("3000000"):
        return "Micro (less than P3000000)"
    if basis <= Decimal("15000000"):
        return "Small ( P3000001 - P15000000)"
    if basis <= Decimal("100000000"):
        return "Medium (P15000001 - P100000000)"
    return "Large (more than P100000000)"


def psic_category(value):
    text = clean(value).lower()
    rules = [
        ("Financial and Insurance Activities", ("bank", "financial", "lending", "pawn", "money", "remittance", "insurance")),
        ("Real Estate Activities", ("real estate", "lessor", "apartment", "rental", "property")),
        ("Wholesale and Retail Trade; Repair of Motor Vehicles and Motorcycles", ("retail", "wholesale", "store", "sari-sari", "pharmacy", "hardware", "motorcycle parts")),
        ("Accommodation and Food Service Activities", ("restaurant", "eatery", "food", "cafe", "hotel", "resort", "lodging", "catering")),
        ("Transportation and Storage", ("transport", "tricycle", "pedicab", "passenger", "cargo", "trucking")),
        ("Manufacturing", ("manufactur", "baking", "bakery", "milling", "printing", "processed")),
        ("Agriculture, Forestry And Fishing", ("agri", "farm", "crop", "forestry", "livestock", "poultry", "fishing")),
        ("Mining and Quarrying", ("mining", "quarry", "sand", "gravel")),
        ("Water Supply; Sewerage, Waste Management And Remediation Activities", ("water", "refilling", "waste", "sewerage")),
        ("Construction", ("construction", "contractor", "building")),
        ("Information and Communication", ("internet", "computer", "telecom", "communication")),
        ("Professional, Scientific and Technical Activities", ("legal", "accounting", "engineering", "consult", "technical")),
        ("Education", ("school", "tutorial", "education", "training")),
        ("Human Health and Social Work Activities", ("clinic", "medical", "dental", "laboratory", "health")),
        ("Arts, Entertainment and Recreation", ("gambling", "cockpit", "amusement", "recreation", "sports")),
    ]
    for category, keywords in rules:
        if any(keyword in text for keyword in keywords):
            return category
    return "Other Service Activities"


def target_sheet(workbook, year):
    if year == 2025 and "Annex A (Jan. to Dec. 2025)" in workbook.sheetnames:
        return workbook["Annex A (Jan. to Dec. 2025)"]
    return next(sheet for sheet in workbook.worksheets if sheet.title.startswith("Annex B"))


def export_report(date_from, date_to, output_dir):
    from openpyxl import load_workbook

    start = datetime.strptime(date_from, "%Y-%m-%d").date()
    end = datetime.strptime(date_to, "%Y-%m-%d").date()
    records = []
    for record in build_payload(5000).get("records", []):
        application_date = parse_date(record.get("application_date"))
        if application_date and start <= application_date <= end and clean(record.get("permit_no")):
            records.append((application_date, record))
    records.sort(key=lambda item: (item[0], clean(item[1].get("permit_no")), clean(item[1].get("business_name"))))

    workbook = load_workbook(TEMPLATE_DIR / "2025-2026_ANNEX-A-B_cmci_report.xlsx")
    grouped = {}
    for application_date, record in records:
        grouped.setdefault(application_date.year, []).append((application_date, record))
    for year, year_rows in grouped.items():
        sheet = target_sheet(workbook, year)
        for row_index in range(7, sheet.max_row + 1):
            for column_index in range(1, 17):
                sheet.cell(row_index, column_index).value = None
        for row_index, (application_date, record) in enumerate(year_rows, start=7):
            values = [
                "Zamboanguita", "Negros Oriental", "REGION VII (CENTRAL VISAYAS)",
                "Third Class Municipality", "Municipality", clean(record.get("business_name")),
                "", clean(record.get("location")) or clean(record.get("barangay")), "",
                clean(record.get("owner_name")),
                psic_category(record.get("business_line") or record.get("business_nature")),
                normalize_business_type(record.get("business_type")),
                capitalization_size(record.get("capital_investment"), record.get("gross_sales")),
                "New" if clean(record.get("application_type")).upper() == "NEW" else "Renewal",
                application_date.year, clean(record.get("permit_no")),
            ]
            for column_index, value in enumerate(values, start=1):
                sheet.cell(row_index, column_index).value = excel_value(value)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    path = save_workbook(workbook, output_path(output_dir, REPORT_NUMBER, date_from, date_to))
    return path, len(records)


if __name__ == "__main__":
    raise SystemExit(run_cli(REPORT_NUMBER, export_report))
