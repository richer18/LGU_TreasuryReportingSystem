from decimal import Decimal

from local_report_export_common import (
    BUSINESS_PERMIT_DIR,
    excel_value,
    output_path,
    run_cli,
    save_workbook,
)
from report_preview_readonly import (
    business_establishment_match_lookup,
    clean_text,
    decimal_value,
    find_business_permit_workbook,
    load_sheet_records,
    load_sheet_records_with_header,
    parse_excel_date,
    tax_on_business_category,
)


REPORT_NUMBER = 33


def build_report_data(date_from, date_to):
    from datetime import datetime

    abstract_path = find_business_permit_workbook("ABSTRACT_OF_GENERAL_COLLECTION-BPLS*.xlsx")
    establishment_path = find_business_permit_workbook("BUSINESS_ESTABLISHMENT-BPLS*.xlsx")
    records = load_sheet_records_with_header(abstract_path, 7)
    establishments = load_sheet_records(establishment_path)
    by_or, by_business_id = business_establishment_match_lookup(establishments)
    start = datetime.strptime(date_from, "%Y-%m-%d").date()
    end = datetime.strptime(date_to, "%Y-%m-%d").date()
    categories = ["Manufacturing", "Distributor", "Retailing", "Banks & Other Financial Int.", "Other Business Tax", "Fines & Penalties"]
    summary = {category: {"business_tax": Decimal("0"), "surcharge": Decimal("0")} for category in categories}
    details = []
    for record in records:
        or_date = parse_excel_date(record.get("O.R. Date"))
        if or_date is None or not start <= or_date <= end:
            continue
        business_tax = decimal_value(record.get("Business Tax"))
        surcharge = decimal_value(record.get("Surcharge"))
        if business_tax == 0 and surcharge == 0:
            continue
        or_number = clean_text(record.get("O.R. Number"))
        business_id = clean_text(record.get("Business Identification Number"))
        establishment = by_or.get(or_number) or by_business_id.get(business_id) or {}
        nature = clean_text(establishment.get("Business Nature"))
        line = clean_text(establishment.get("Business Line"))
        category = tax_on_business_category(nature, line)
        summary[category]["business_tax"] += business_tax
        summary["Fines & Penalties"]["surcharge"] += surcharge
        details.append([
            or_date, clean_text(record.get("Date Paid")), or_number,
            clean_text(record.get("Transaction Type")), business_id,
            clean_text(record.get("Business Name")), clean_text(record.get("Barangay Name")),
            nature, line, category, business_tax, surcharge,
            decimal_value(record.get("Amount Paid")),
            "OR Number" if or_number in by_or else "Business ID" if business_id in by_business_id else "Unmatched",
        ])
    return {
        "summary": [[category, summary[category]["business_tax"], summary[category]["surcharge"], summary[category]["business_tax"] + summary[category]["surcharge"]] for category in categories],
        "details": details,
        "sources": (str(abstract_path), str(establishment_path)),
    }


def style_sheet(sheet):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="B7C9D6")
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)


def export_report(date_from, date_to, output_dir):
    from openpyxl import Workbook

    data = build_report_data(date_from, date_to)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    detail = workbook.create_sheet("Detail")
    notes = workbook.create_sheet("Notes")
    summary.append(["Category", "Business Tax", "Fines & Penalties / Surcharge", "Total"])
    for row in data["summary"]:
        summary.append([excel_value(value) for value in row])
    total_row = summary.max_row + 1
    summary.append(["TOTAL", f"=SUM(B2:B{total_row - 1})", f"=SUM(C2:C{total_row - 1})", f"=SUM(D2:D{total_row - 1})"])
    detail.append(["O.R. Date", "Date Paid", "O.R. Number", "Transaction Type", "Business ID", "Business Name", "Barangay", "Business Nature", "Business Line", "Tax Category", "Business Tax", "Surcharge", "Amount Paid", "Match Basis"])
    for row in data["details"]:
        detail.append([excel_value(value) for value in row])
    notes.append(["Report", "33. Tax on Business Summary from BPLS Business Tax"])
    notes.append(["Period", f"{date_from} to {date_to}"])
    notes.append(["General Collection Source", data["sources"][0]])
    notes.append(["Business Establishment Source", data["sources"][1]])
    for sheet in (summary, detail, notes):
        style_sheet(sheet)
        sheet.freeze_panes = "A2"
    detail.sheet_state = "hidden"
    notes.sheet_state = "hidden"
    path = save_workbook(workbook, output_path(output_dir, REPORT_NUMBER, date_from, date_to))
    return path, len(data["details"])


if __name__ == "__main__":
    raise SystemExit(run_cli(REPORT_NUMBER, export_report))
